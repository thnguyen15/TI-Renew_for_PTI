import streamlit as st
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
import io
from datetime import datetime

# ─── Mapping cột từ file input (0-indexed) ───────────────────────────────────
COL = {
    "dai_ly":       1,   # Đại lý
    "chi_nhanh":    4,   # Chi nhánh
    "so_gcn":       5,   # Số GCN
    "hieu_luc_den": 10,  # Hiệu lực đến
    "ten_kh":       16,  # Tên khách hàng
    "bien_xe":      24,  # Biển xe
    "ten_goi":      34,  # Tên gói bảo hiểm
    "phi_sau_vat":  36,  # Phí sau VAT (phí năm trước)
    "so_vu":        37,  # Số vụ tổn thất
    "so_tien_bt":   38,  # Số tiền bồi thường
    # Gói Bạc (col 41-45)
    "bac_ti_le":    41,
    "bac_phi_truoc":42,
    "bac_giam_phi": 43,
    "bac_ti_le_giam":44,
    "bac_phi_sau":  45,
    # Gói Vàng (col 46-50)
    "vang_ti_le":   46,
    "vang_phi_truoc":47,
    "vang_giam_phi":48,
    "vang_ti_le_giam":49,
    "vang_phi_sau": 50,
    # Gói Bạch Kim (col 51-55)
    "bk_ti_le":     51,
    "bk_phi_truoc": 52,
    "bk_giam_phi":  53,
    "bk_ti_le_giam":54,
    "bk_phi_sau":   55,
    # Giá trị xe năm nay
    "gtxe_nam_nay": 56,
}

# Màu sắc
COLOR_HEADER_MAIN = "1F3864"   # Xanh đậm
COLOR_HEADER_SUB  = "2E75B6"   # Xanh nhạt hơn
COLOR_BAC         = "D9E1F2"   # Xanh nhạt - Gói Bạc
COLOR_VANG        = "FFF2CC"   # Vàng nhạt - Gói Vàng
COLOR_BK          = "E2EFDA"   # Xanh lá nhạt - Gói Bạch Kim
COLOR_ROW_ALT     = "F5F5F5"   # Xám nhạt cho dòng chẵn

FONT_WHITE   = Font(name="Arial", bold=True, color="FFFFFF", size=10)
FONT_HEADER  = Font(name="Arial", bold=True, color="FFFFFF", size=9)
FONT_DATA    = Font(name="Arial", size=9)
FONT_BOLD    = Font(name="Arial", bold=True, size=9)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
ALIGN_RIGHT  = Alignment(horizontal="right",  vertical="center")

THIN = Side(style="thin", color="BFBFBF")
BORDER_THIN = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FMT_NUMBER = '#,##0'
FMT_PCT    = '0.00"%"'


def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def apply_cell(ws, row, col, value=None, font=None, fill_=None,
               align=None, border=None, fmt=None):
    c = ws.cell(row=row, column=col, value=value)
    if font:   c.font      = font
    if fill_:  c.fill      = fill_
    if align:  c.alignment = align
    if border: c.border    = border
    if fmt:    c.number_format = fmt
    return c


def build_output_xlsx(df_input):
    """
    Nhận DataFrame từ file input (header=None, row 0-1 là title, row 3-4 là header,
    row 5+ là data). Trả về bytes của file .xlsx output.
    """
    # Lấy data rows (từ row index 5 trở đi)
    data = df_input.iloc[5:].reset_index(drop=True)
    data.columns = range(len(data.columns))

    # Lấy danh sách đại lý unique (giữ thứ tự xuất hiện)
    dai_ly_list = data[COL["dai_ly"]].dropna().unique().tolist()

    wb = Workbook()
    wb.remove(wb.active)  # xóa sheet mặc định

    for dai_ly in dai_ly_list:
        df_dl = data[data[COL["dai_ly"]] == dai_ly].reset_index(drop=True)

        # Tên sheet: lấy tối đa 31 ký tự, bỏ ký tự không hợp lệ
        sheet_name = str(dai_ly)[:31].replace("/", "-").replace("\\", "-")
        ws = wb.create_sheet(title=sheet_name)

        # ── ROW 1: Tiêu đề bảng ─────────────────────────────────────────────
        ws.merge_cells("A1:O1")
        c = ws["A1"]
        c.value   = f"BẢNG BÁO PHÍ BẢO HIỂM - {str(dai_ly).upper()}"
        c.font    = Font(name="Arial", bold=True, size=13, color="FFFFFF")
        c.fill    = fill(COLOR_HEADER_MAIN)
        c.alignment = ALIGN_CENTER

        # ── ROW 2-3: Header hai dòng ─────────────────────────────────────────
        # Cột đơn (merge 2 dòng):
        single_cols = [
            (1,  "STT"),
            (2,  "Số GCN"),
            (3,  "Hiệu lực đến"),
            (4,  "Chi nhánh"),
            (5,  "Tên khách hàng"),
            (6,  "Biển xe"),
            (7,  "Tên gói BH"),
            (8,  "Phí năm trước"),
        ]
        for col_idx, label in single_cols:
            ws.merge_cells(start_row=2, start_column=col_idx,
                           end_row=3,   end_column=col_idx)
            apply_cell(ws, 2, col_idx, label, FONT_HEADER, fill(COLOR_HEADER_SUB),
                       ALIGN_CENTER, BORDER_THIN)

        # Nhóm "Tỉ lệ bồi thường năm trước" (col 9-10, merge ngang)
        ws.merge_cells(start_row=2, start_column=9, end_row=2, end_column=10)
        apply_cell(ws, 2, 9, "Tỉ lệ bồi thường năm trước",
                   FONT_HEADER, fill(COLOR_HEADER_SUB), ALIGN_CENTER, BORDER_THIN)
        apply_cell(ws, 3, 9,  "Số vụ TT",     FONT_HEADER, fill(COLOR_HEADER_SUB), ALIGN_CENTER, BORDER_THIN)
        apply_cell(ws, 3, 10, "Số tiền BT",   FONT_HEADER, fill(COLOR_HEADER_SUB), ALIGN_CENTER, BORDER_THIN)

        # Nhóm "Phí Tái Tục Năm Nay" (col 11-15, merge ngang)
        ws.merge_cells(start_row=2, start_column=11, end_row=2, end_column=15)
        apply_cell(ws, 2, 11, "Phí Tái Tục Năm Nay",
                   FONT_HEADER, fill(COLOR_HEADER_SUB), ALIGN_CENTER, BORDER_THIN)
        sub_headers = ["Giá trị xe", "Tỉ lệ phí (%)", "Tỷ lệ giảm (%)", "Phí năm nay", "Phí sau giảm"]
        for i, h in enumerate(sub_headers):
            apply_cell(ws, 3, 11 + i, h, FONT_HEADER, fill(COLOR_HEADER_SUB),
                       ALIGN_CENTER, BORDER_THIN)

        # ── ROW 4+: Data ──────────────────────────────────────────────────────
        def safe(val):
            if pd.isna(val): return None
            return val

        def fmt_date(val):
            if pd.isna(val): return None
            if isinstance(val, datetime): return val.strftime("%d/%m/%Y")
            return str(val)

        for r_idx, (_, row) in enumerate(df_dl.iterrows()):
            excel_row = 4 + r_idx
            is_even = r_idx % 2 == 1
            row_fill = fill(COLOR_ROW_ALT) if is_even else None

            # Xác định gói BH để lấy đúng cột phí
            goi = str(safe(row[COL["ten_goi"]]) or "").strip().lower()
            if "vàng" in goi or "vang" in goi:
                ti_le_phi  = safe(row[COL["vang_ti_le"]])
                phi_truoc  = safe(row[COL["vang_phi_truoc"]])
                ti_le_giam = safe(row[COL["vang_ti_le_giam"]])
                phi_sau    = safe(row[COL["vang_phi_sau"]])
            elif "bạch kim" in goi or "bach kim" in goi or "bạchkim" in goi:
                ti_le_phi  = safe(row[COL["bk_ti_le"]])
                phi_truoc  = safe(row[COL["bk_phi_truoc"]])
                ti_le_giam = safe(row[COL["bk_ti_le_giam"]])
                phi_sau    = safe(row[COL["bk_phi_sau"]])
            else:  # Bạc (default)
                ti_le_phi  = safe(row[COL["bac_ti_le"]])
                phi_truoc  = safe(row[COL["bac_phi_truoc"]])
                ti_le_giam = safe(row[COL["bac_ti_le_giam"]])
                phi_sau    = safe(row[COL["bac_phi_sau"]])

            cells_data = [
                (1,  r_idx + 1,                        FONT_DATA, ALIGN_CENTER, None),
                (2,  safe(row[COL["so_gcn"]]),          FONT_DATA, ALIGN_LEFT,   None),
                (3,  fmt_date(row[COL["hieu_luc_den"]]),FONT_DATA, ALIGN_CENTER, None),
                (4,  safe(row[COL["chi_nhanh"]]),       FONT_DATA, ALIGN_LEFT,   None),
                (5,  safe(row[COL["ten_kh"]]),          FONT_DATA, ALIGN_LEFT,   None),
                (6,  safe(row[COL["bien_xe"]]),         FONT_DATA, ALIGN_CENTER, None),
                (7,  safe(row[COL["ten_goi"]]),         FONT_DATA, ALIGN_CENTER, None),
                (8,  safe(row[COL["phi_sau_vat"]]),     FONT_DATA, ALIGN_RIGHT,  FMT_NUMBER),
                (9,  safe(row[COL["so_vu"]]),           FONT_DATA, ALIGN_CENTER, None),
                (10, safe(row[COL["so_tien_bt"]]),      FONT_DATA, ALIGN_RIGHT,  FMT_NUMBER),
                (11, safe(row[COL["gtxe_nam_nay"]]),    FONT_DATA, ALIGN_RIGHT,  FMT_NUMBER),
                (12, ti_le_phi,                         FONT_DATA, ALIGN_CENTER, '0.00'),
                (13, ti_le_giam,                        FONT_DATA, ALIGN_CENTER, '0.00'),
                (14, phi_truoc,                         FONT_DATA, ALIGN_RIGHT,  FMT_NUMBER),
                (15, phi_sau,                           FONT_BOLD, ALIGN_RIGHT,  FMT_NUMBER),
            ]

            for col_idx, value, font, align, fmt in cells_data:
                c = apply_cell(ws, excel_row, col_idx, value, font,
                               row_fill, align, BORDER_THIN, fmt)

        # ── Độ rộng cột ──────────────────────────────────────────────────────
        col_widths = [5, 24, 14, 16, 22, 12, 12, 14, 10, 18, 16, 12, 14, 16, 16]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        # Chiều cao dòng
        ws.row_dimensions[1].height = 28
        ws.row_dimensions[2].height = 20
        ws.row_dimensions[3].height = 35

        # Freeze panes
        ws.freeze_panes = "A4"

    # ── Ghi ra bytes ─────────────────────────────────────────────────────────
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# ════════════════════════════════════════════════════════════════════════════
#  STREAMLIT UI
# ════════════════════════════════════════════════════════════════════════════

st.set_page_config(
    page_title="Báo Phí Tái Tục",
    page_icon="🚗",
    layout="centered"
)

st.title("🚗 Hệ thống báo phí tái tục")
st.caption("Toyota Tsusho Insurance Service Vietnam")
st.markdown("---")

st.markdown("""
**Hướng dẫn sử dụng:**
1. Tải file báo cáo dự báo tái tục (`.xlsx`) lên bên dưới
2. Hệ thống sẽ tự động xử lý và tạo bảng báo phí cho từng đại lý
3. Tải file kết quả về và gửi cho đại lý
""")

uploaded_file = st.file_uploader(
    "📂 Chọn file báo cáo dự báo tái tục",
    type=["xlsx"],
    help="File xuất từ hệ thống, định dạng .xlsx"
)

if uploaded_file is not None:
    st.success(f"✅ Đã nhận file: **{uploaded_file.name}**")

    with st.spinner("Đang xử lý dữ liệu..."):
        try:
            df_input = pd.read_excel(uploaded_file, header=None, sheet_name=0)

            # Kiểm tra sơ bộ cấu trúc file
            if df_input.shape[1] < 45:
                st.error("❌ File không đúng định dạng. Vui lòng kiểm tra lại.")
                st.stop()

            data_rows = df_input.iloc[5:]
            dai_ly_col = data_rows.iloc[:, COL["dai_ly"]].dropna()
            n_contracts = len(data_rows.dropna(subset=[data_rows.columns[COL["dai_ly"]]]))
            n_dealers   = len(dai_ly_col.unique())

            st.info(f"📊 Tìm thấy **{n_contracts} hợp đồng** / **{n_dealers} đại lý**")

            # Hiển thị danh sách đại lý
            dealer_names = sorted(dai_ly_col.unique().tolist())
            with st.expander("Xem danh sách đại lý"):
                for i, name in enumerate(dealer_names, 1):
                    count = (dai_ly_col == name).sum()
                    st.write(f"{i}. **{name}** — {count} hợp đồng")

            # Build output
            output_bytes = build_output_xlsx(df_input)

            # Tên file output
            today_str = datetime.now().strftime("%Y%m%d")
            output_filename = f"Bang_bao_phi_tai_tuc_{today_str}.xlsx"

            st.success("🎉 Xử lý thành công!")
            st.download_button(
                label="⬇️ Tải file kết quả xuống",
                data=output_bytes,
                file_name=output_filename,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary"
            )

        except Exception as e:
            st.error(f"❌ Lỗi khi xử lý file: {str(e)}")
            st.exception(e)

st.markdown("---")
st.caption("v1.0 · TTISV Internal Tool")
